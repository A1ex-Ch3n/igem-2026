#!/usr/bin/env python
"""Apply the MRA00793-26 revision into the ORIGINAL submitted .docx, preserving its formatting.

PROVENANCE ONLY. Run once on 2026-08-09 to produce
`documents/2026-08-08_mra_revision/manuscript_revised.docx` from
`documents/2026-08-08_mra_revision/manuscript_submitted.docx` (the file as submitted), applying the change
list in `documents/2026-08-08_mra_revision/response_to_reviewers.md`.

It copies the submitted file and edits it in place at the XML level, so styles, fonts,
margins, line numbering and table borders are the journal file's own, not a
reconstruction. Run properties are cloned from each paragraph's first run before the
text is replaced, then italics / subscripts / highlighting are re-applied.

**The .docx is authoritative the moment it is opened in Word.** Do not re-run this to
"regenerate" the manuscript after any hand edit -- it starts again from the submitted
file and will silently discard that edit.

Asserts the MRA limits before writing: abstract <= 50 words, Announcement <= 500.
"""
import re, copy, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/Users/alexy/Desktop/iGEM 26/documents/2026-08-08_mra_revision/manuscript_submitted.docx"
OUT = "/Users/alexy/Desktop/iGEM 26/documents/2026-08-08_mra_revision/manuscript_revised.docx"

TOKEN = re.compile(r"(\|i\||\|/i\||\|s\||\|/s\||\|h\||\|/h\|)")

def strip(t): return TOKEN.sub("", t)

def set_para(p, markup, template_run=None):
    """Replace a paragraph's text, cloning run properties from the paragraph's own first run."""
    runs = p.runs
    tmpl = template_run if template_run is not None else (runs[0] if runs else None)
    tmpl_rPr = copy.deepcopy(tmpl._element.rPr) if (tmpl is not None and tmpl._element.rPr is not None) else None
    for r in list(runs):
        r._element.getparent().remove(r._element)
    ital = sub = hl = False
    for tok in TOKEN.split(markup):
        if tok == "|i|": ital = True; continue
        if tok == "|/i|": ital = False; continue
        if tok == "|s|": sub = True; continue
        if tok == "|/s|": sub = False; continue
        if tok == "|h|": hl = True; continue
        if tok == "|/h|": hl = False; continue
        if not tok: continue
        r = p.add_run(tok)
        if tmpl_rPr is not None:
            r._element.remove(r._element.get_or_add_rPr())
            r._element.insert(0, copy.deepcopy(tmpl_rPr))
        r.italic = ital or None
        if sub: r.font.subscript = True
        if hl:
            h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow")
            r._element.get_or_add_rPr().append(h)
    return p

doc = None  # opened after the word-limit guard below, so a failed run leaves no output file
P = None

# ---------------------------------------------------------------- body text
ABSTRACT = ("We report the complete genome sequences of three |i|Pseudomonas|/i| isolates recovered from topsoil "
            "at the California Botanic Garden. Two isolates share ~99.4% ANI, enabling investigation of "
            "intraspecies microvariation; the third represents a co-occurring distinct species, together "
            "capturing species- and strain-level genomic diversity within a natural soil community.")

INTRO = ("High-fidelity, complete genomes of closely related environmental isolates, such as those of the abundant "
         "and diverse |i|Pseudomonas|/i| genus[1], provide a valuable resource for studying micro diversity at full "
         "genome resolution. Here, we report the complete genome sequences of three |i|Pseudomonas|/i| isolates "
         "recovered from soil samples at the California Botanic Garden (Claremont, CA).")

ISOLATION = ("Soil samples were collected with a T-Handle soil probe (Lichtund); collection sites are detailed in "
             "Table 1. Soil was transported to Keck Graduate Institute (KGI) and stored at 4°C until further "
             "processing. Bacterial enrichments were performed using 5 g of soil in 25 mL of |i|Pseudomonas|/i| "
             "Minimal Media (2 mM MgSO|s|4|/s|, 0.31 mM FeCl|s|3|/s|, 0.03 mM CaCl|s|2|/s|, 15 mM NaH|s|2|/s|PO|s|4|/s|, "
             "15 mM K|s|2|/s|HPO|s|4|/s|, pH 6.8) amended with 5 drops each of 10% Na-succinate and 20% NH|s|4|/s|Cl as "
             "carbon and nitrogen sources [2]. After two days, enrichments were plated with a 1 uL loop on "
             "|i|Pseudomonas|/i| Minimal Agar (1.5% agar); after a further two days colonies were picked onto "
             "|i|Pseudomonas|/i| Isolation Agar (PIA, Criterion No:C6681, with 10 mL glycerol per 500 mL). Single "
             "colonies were grown overnight in lysogeny broth (LB; 25 g/L). Next morning, cultures reaching an "
             "optical density (OD600) > 0.6 were archived as 25% glycerol stocks at -80°C; all 28 isolates "
             "archived in this study met this threshold. Five isolates spanning three collection sites and both "
             "colony types recovered on PIA were selected for whole-genome sequencing, and collapsed by ANI into "
             "the three strains reported here.")

SEQ = ("For genome sequencing, colonies were revived from stocks by single-colony streak on LB agar (1% agar) and "
       "grown in 10 mL LB (2.5%). Overnight cultures were then pelleted at 3000 rpm for 15 minutes, washed twice "
       "with 5 mL PBS and once with 1.5 mL PBS. DNA was extracted using the QuickDNA Kit (Zymo Research "
       "Corporation), following manufacturer’s instructions. DNA was quantified using Portlite Fluorimetric DNA "
       "Quantification Kit (AAT Bioquest) and renormalized to 50 ng/uL. 20 uL were sent out for sequencing at "
       "Plasmidsaurus (Eugene, OR) on an Oxford Nanopore PromethION using R10.4.1 flow cells, with "
       "|h|amplification-free v14 chemistry in which input DNA undergoes minimal sequence-independent fragmentation "
       "by tagmentation|/h|; no size selection was performed.")

DEFAULTS = "Default parameters were used for all software unless otherwise specified."

BIOINF = ("Basecalling was performed with Dorado using the super-accuracy (sup) basecalling model (v4.3.0 [3]) with "
          "default Q10 quality filtering, and the poorest 5% of reads were removed with Filtlong (v0.2.1 [4]). A de "
          "novo assembly was generated with Autocycler (v0.2.1[5]) from input assemblies produced by Flye (v2.9.6 "
          "[6]), hifiasm ([7]ONT mode), and Plassembler (v1.8.0 [8]). The circular assembly graph was reoriented "
          "using dnaapler (v1.2.0[9]). Genome completeness and contamination were assessed with CheckM (v1.2.2[10]). "
          "Isolates were taxonomically classified with sourmash (v4.9.4 [11]) against a GTDB database, which "
          "assigned all isolates to the genus |i|Pseudomonas|/i|, two with an average nucleotide identity (ANI) of "
          "99.5–99.75% to |i|Pseudomonas|/i| sp013179515 (GCF_013179515.1) and one to |i|Pseudomonas capeferrum|/i| "
          "(GCF_000731675.1, ANI 98.66 %). Genomes were annotated with Bakta (v1.12.0 [12]). Sequencing and "
          "assembly statistics are summarized in Table 1.")

DATA = ("The complete genomes of our isolates and the raw sequencing data are available under BioProject "
        "PRJNA1482863 and PRJNA1483221 respectively. Individual accession numbers are listed in Table 1.")

REF2 = ("K. Sandman and C. Ecker, “|i|Pseudomonas|/i| Isolation and Identification: An Introduction to the "
        "Challenges of Polyphasic Taxonomy,” ")

CAPTION_OLD_TAIL = None  # handled below

n_abs = len(strip(ABSTRACT).split())
n_body = sum(len(strip(x).split()) for x in (INTRO, ISOLATION, SEQ, DEFAULTS, BIOINF))
print(f"Abstract {n_abs}/50   Announcement {n_body}/500")
assert n_abs <= 50 and n_body <= 500, "over word limit"

# only now touch the filesystem: a tripped guard must not leave a half-done OUT behind
shutil.copyfile(SRC, OUT)
doc = Document(OUT)
P = doc.paragraphs

set_para(P[8], ABSTRACT)
set_para(P[10], INTRO)
set_para(P[11], ISOLATION)
set_para(P[12], SEQ)
set_para(P[13], DEFAULTS)
set_para(P[14], BIOINF)
set_para(P[16], DATA)

def italicise_word(par, word, marker=None):
    """Split the run containing `marker or word` so `word` becomes its own italic run."""
    key = marker or word
    from docx.text.run import Run
    for r in par.runs:
        if key not in r.text: continue
        pre, post = r.text.split(word, 1)
        el = r._element
        mid = copy.deepcopy(el); tail = copy.deepcopy(el)
        r.text = pre
        el.addnext(tail); el.addnext(mid)
        Run(mid, par).text = word
        Run(mid, par).italic = True
        rt = Run(tail, par); rt.text = post; rt.italic = None
        return True
    return False

assert italicise_word(P[22], "Pseudomonas", "\u201cPseudomonas Isolation and Identification"), "ref [2] genus not found"
assert italicise_word(P[42], "Pseudomonas", "three Pseudomonas isolates"), "caption genus not found"

# ---------------------------------------------------------------- Table 1
t = doc.tables[0]

def cell_set(cell, markup):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para(p, markup)

def row_after(idx):
    """Clone row idx (keeps borders/shading/widths) and insert directly after it."""
    tr = copy.deepcopy(t.rows[idx]._tr)
    t.rows[idx]._tr.addnext(tr)
    return t.rows[idx + 1]

def find_row(label):
    for i, row in enumerate(t.rows):
        if row.cells[0].text.strip() == label: return i
    raise KeyError(label)

# 1. delete the redundant first data row (R2.3)
i = find_row("Isolate (internal name)")
t.rows[i]._tr.getparent().remove(t.rows[i]._tr)

# 2. fix the "Represents" row (R1.6)
i = find_row("Represents (≈100%-ANI clones)")
cell_set(t.rows[i].cells[0], "Additional ≈100%-ANI isolate represented")
for c, v in zip(t.rows[i].cells[1:], ["10-10-2 (99.997% ANI)", "9-6-2 (100.000% ANI)", "none"]):
    cell_set(c, v)

# 3. BioSample row relabelled, and a new assembly-accession row after it (R1.7)
i = find_row("Genome accession (GenBank)")
cell_set(t.rows[i].cells[0], "BioSample")
new = row_after(i)
cell_set(new.cells[0], "Genome assembly (GenBank)")
for c, v in zip(new.cells[1:], ["GCA_059295255.1", "GCA_059295275.1", "GCA_059295235.1"]):
    cell_set(c, v)

# 4. adjacent plants italicised (R1.8 / R2.1)
i = find_row("Adjacent plant")
for c, v in zip(t.rows[i].cells[1:], ["|i|Chilopsis linearis|/i|", "|i|Aesculus californica|/i|", "|i|Chilopsis linearis|/i|"]):
    cell_set(c, v)

# 5. platform now names instrument + flow cell (R1.5 / R2.2)
i = find_row("Sequencing platform")
for c in t.rows[i].cells[1:]:
    cell_set(c, "Oxford Nanopore PromethION, R10.4.1")

# 6. coverage / read rows split into run vs assembly (R2.2)
i = find_row("Sequencing depth (×)")
cell_set(t.rows[i].cells[0], "Final assembly coverage (×)")

i = find_row("Total reads")
cell_set(t.rows[i].cells[0], "Total raw reads")
cell_set(t.rows[i].cells[3], "|h|[PENDING — see query]|/h|")

bases = row_after(i)
cell_set(bases.cells[0], "Total bases sequenced")
for c, v in zip(bases.cells[1:], ["3,026 Mb", "1,935 Mb", "|h|[PENDING — see query]|/h|"]):
    cell_set(c, v)

rawcov = row_after(i + 1)
cell_set(rawcov.cells[0], "Raw coverage (×)")
for c, v in zip(rawcov.cells[1:], ["510", "313", "|h|[PENDING — see query]|/h|"]):
    cell_set(c, v)

# 7. closest species italicised (R1.8 / R2.1)
i = find_row("Closest species (GTDB)")
for c, v in zip(t.rows[i].cells[1:],
                ["|i|Pseudomonas_E sp013179515|/i|", "|i|Pseudomonas_E sp013179515|/i|", "|i|Pseudomonas_E capeferrum|/i|"]):
    cell_set(c, v)

# ---------------------------------------------------------------- working note
NOTE = ("|h|AUTHOR WORKING NOTE — DELETE THIS PARAGRAPH BEFORE SUBMISSION. Two items are still open and are "
        "highlighted in yellow below: (1) the library-preparation description, rewritten from Plasmidsaurus’s "
        "technical documentation and awaiting confirmation; (2) the three 8-10-2 read-count cells in Table 1, which "
        "cannot be filled until the discrepancy with SRR39379360 is resolved. Everything else is final.|/h|")
first = doc.paragraphs[0]
note_p = copy.deepcopy(first._element)
first._element.addprevious(note_p)
set_para(doc.paragraphs[0], NOTE)

doc.save(OUT)
print("wrote", OUT)

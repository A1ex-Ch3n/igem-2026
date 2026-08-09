#!/usr/bin/env python
"""Render the submission-ready half of the response letter to .docx.

Source of truth is the markdown; this only re-renders it, so edit the .md and
re-run rather than editing the .docx. Stops at the "NOT FOR SUBMISSION" divider
so internal notes cannot be uploaded to the journal by accident.
"""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/Users/alexy/Desktop/iGEM 26/documents/2026-08-08_mra_revision/response_to_reviewers.md"
OUT = "/Users/alexy/Desktop/iGEM 26/documents/2026-08-08_mra_revision/response_to_reviewers.docx"
STOP = "# NOT FOR SUBMISSION"

md = open(SRC, encoding="utf8").read()
cut = md.find(STOP)
assert cut > 0, "divider not found — refusing to render the whole file"
md = md[:cut]
md = re.sub(r"\n-{3,}\n-{3,}\n", "\n", md)          # the double rule before the divider
lines = md.split("\n")

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`]+`)")

def emit(par, text, base_bold=False, base_italic=False, size=11, highlight=False):
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold, ital, mono = base_bold, base_italic, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, bold = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*"):
            tok, ital = tok[1:-1], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, mono = tok[1:-1], True
        r = par.add_run(tok)
        r.font.size = Pt(size)
        r.font.name = "Consolas" if mono else "Calibri"
        r.bold, r.italic = bold or None, ital or None
        if highlight:
            h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow")
            r._element.get_or_add_rPr().append(h)
    return par

def rule(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), "BFBFBF")
    pbdr.append(b); pPr.append(pbdr)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(8); st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(1.0)

buf = []
def flush():
    """Emit the accumulated paragraph lines as one justified paragraph."""
    if not buf: return
    text = " ".join(buf).strip(); buf.clear()
    if not text: return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    emit(p, text)

for raw in lines:
    line = raw.rstrip()
    if not line.strip():
        flush(); continue
    if line.startswith("# "):
        flush()
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        emit(p, line[2:], base_bold=True, size=16)
        continue
    if line.startswith("## "):
        flush()
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
        emit(p, line[3:], base_bold=True, size=13)
        continue
    if re.fullmatch(r"-{3,}", line.strip()):
        flush()
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10); rule(p)
        continue
    if line.lstrip().startswith(">"):
        flush()
        body = line.lstrip()[1:].strip()
        if not body: continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.3)
        emit(p, body, highlight=body.startswith("🛑"))
        continue
    # a line that is entirely bold is its own line (date / byline), not a continuation
    if re.fullmatch(r"\*\*.+\*\*", line.strip()):
        flush()
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        emit(p, line.strip())
        continue
    buf.append(line.strip())
flush()

doc.save(OUT)

d2 = Document(OUT)
print("wrote", OUT)
print("paragraphs:", len(d2.paragraphs))
print("contains internal notes:", "NOT FOR SUBMISSION" in " ".join(p.text for p in d2.paragraphs))
